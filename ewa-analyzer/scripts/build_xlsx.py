#!/usr/bin/env python3
"""Build the EWA analysis workbook from findings.json.   python build_xlsx.py findings.json out.xlsx
Needs: pip install openpyxl"""
import sys
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
except ImportError:
    sys.exit("openpyxl is not installed. Run: pip install openpyxl")
from common import load_findings, FindingsError, PRIORITIES, est_lines, risk_areas

NAVY, BLUE, GREY = "1F3864", "2E5395", "595959"
PCOL = {"Critical": "C00000", "High": "ED7D31", "Medium": "FFC000", "Low": "70AD47"}
TREND = {"improved": "388E3C", "worse": "C62828", "slightly worse": "F57C00", "unchanged": "595959"}
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(sz=10, b=False, color=None, i=False):
    return Font(name="Arial", size=sz, bold=b, italic=i, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def header_row(ws, row, labels):
    for c, t in enumerate(labels, 1):
        x = ws.cell(row, c, t)
        x.font, x.fill, x.border = font(10, True, "FFFFFF"), fill(NAVY), BOX
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 20


def body_cell(ws, r, c, v, bold=False, center=False, color=None):
    x = ws.cell(r, c, v)
    x.font = font(10, bold, color)
    x.border = BOX
    x.alignment = Alignment(horizontal="center" if center else "left", vertical="center" if center else "top", wrap_text=True)
    return x


def row_height(values, widths, minimum=15):
    lines = 1
    for v, w in zip(values, widths):
        lines = max(lines, est_lines(v if v is not None else "", w * 1.05))
    return max(minimum, lines * 13 + 4)


def finish(ws, landscape=True, title_row=3):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if title_row:
        ws.print_title_rows = "%d:%d" % (title_row, title_row)


def exec_summary(wb, d):
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", (30, 55, 30)):
        ws.column_dimensions[col].width = w
    m, o = d["meta"], d["overall"]
    ws.merge_cells("A1:C1"); ws["A1"] = "SAP EarlyWatch Alert — Detailed Analysis"
    ws["A1"].font = font(16, True, NAVY); ws.row_dimensions[1].height = 19.5
    sub = "System: %s%s | %s" % (m["system_id"], (" (%s)" % m["product"]) if m.get("product") else "", m.get("landscape", ""))
    ws.merge_cells("A2:C2"); ws["A2"] = sub.rstrip(" |")
    ws.merge_cells("A3:C3"); ws["A3"] = "Prepared by: %s  |  EWA report dated %s  |  Period: %s" % (m["prepared_by"], m["report_date"], m["period"])
    for a in ("A2", "A3"):
        ws[a].font = font(10, False, GREY)
    ws.merge_cells("A5:C5"); ws["A5"] = "Overall System Rating"
    ws["A5"].font, ws["A5"].fill = font(12, True, "FFFFFF"), fill(BLUE)
    for c in "BC":
        ws[c + "5"].fill = fill(BLUE)
    rows = [("Overall EWA Rating", o.get("rating_text") or o["label"]), ("Decisive Red Alert", o["decisive_alert"])] + [tuple(x) for x in d.get("system_facts", [])]
    r = 6
    for k, v in rows:
        ws.cell(r, 1, k).font = font(10, True)
        ws.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        x = ws.cell(r, 2, v); x.font = font(10); x.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = row_height([v], [85], 15)
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, "Action Items by Priority").font = font(12, True, "FFFFFF")
    for c in range(1, 4):
        ws.cell(r, c).fill = fill(BLUE)
    r += 1
    header_row(ws, r, ["Priority", "Count", "Primary Risk Area"])
    ra = risk_areas(d)
    for pr in PRIORITIES:
        r += 1
        x = ws.cell(r, 1, pr); x.font, x.fill, x.border = font(10, True, "FFFFFF"), fill(PCOL[pr]), BOX
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        body_cell(ws, r, 2, "=COUNTIF('Action Items'!B:B,\"%s\")" % pr, center=True)
        body_cell(ws, r, 3, ra[pr])
        ws.row_dimensions[r].height = row_height([ra[pr]], [32], 30)
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, "Top Immediate Actions").font = font(12, True, "FFFFFF")
    for c in range(1, 4):
        ws.cell(r, c).fill = fill(BLUE)
    for t in d.get("top_actions", []):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        x = ws.cell(r, 1, t); x.font = font(10); x.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = row_height([t], [125], 15)
    finish(ws, landscape=False, title_row=None)
    return ws


def action_items(wb, d):
    ws = wb.create_sheet("Action Items")
    widths = [4, 10, 16, 30, 38, 30, 42, 12, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"] = "%s EarlyWatch Alert — Full Action Item Register (%s)" % (d["meta"]["system_id"], d["meta"]["period"])
    ws["A1"].font = font(16, True, NAVY); ws.row_dimensions[1].height = 20
    header_row(ws, 3, ["#", "Priority", "Category", "Finding", "Technical Detail", "Business Impact", "Recommendation", "Owner", "Status"])
    order = {p: i for i, p in enumerate(PRIORITIES)}
    items = sorted(d["action_items"], key=lambda a: order[a["priority"]])  # stable: keeps the analyst's order within a priority
    r = 3
    for n, a in enumerate(items, 1):
        r += 1
        vals = [n, a["priority"], a["category"], a["finding"], a.get("technical", ""), a.get("impact", ""), a["recommendation"], a.get("owner", ""), a.get("status", "Open")]
        for c, v in enumerate(vals, 1):
            if c == 2:
                x = body_cell(ws, r, c, v, True, True, "FFFFFF"); x.fill = fill(PCOL[v])
            elif c == 4:
                body_cell(ws, r, c, v, True)
            else:
                body_cell(ws, r, c, v, False, c in (1, 8, 9))
        ws.row_dimensions[r].height = row_height(vals, widths)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:I%d" % r
    finish(ws)
    return ws


def detail_sheet(wb, spec):
    ws = wb.create_sheet(spec["name"][:31])
    cols = spec["columns"]
    widths = spec.get("widths") or [max(14, min(60, int(100 / len(cols))))] * len(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"] = spec.get("title", spec["name"]); ws["A1"].font = font(14, True, NAVY)
    ws.row_dimensions[1].height = 20
    header_row(ws, 3, cols)
    r = 3
    for row in spec["rows"]:
        r += 1
        row = list(row) + [""] * (len(cols) - len(row))
        for c, v in enumerate(row[:len(cols)], 1):
            x = body_cell(ws, r, c, v)
            col = spec.get("trend_col")
            if col and c == col + 1:
                x.font = font(10, True, TREND.get(str(v).lower(), NAVY if v else None))
                x.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = row_height(row, widths)
    ws.freeze_panes = "A4"
    if spec.get("hidden"):
        ws.sheet_state = "hidden"
    finish(ws)
    return ws


def change_sheet(wb, d):
    prev = d["previous"]
    cur = d["meta"].get("short_period") or d["meta"]["period"]
    spec = {
        "name": "Change vs %s" % prev.get("short_label", "previous")[:18],
        "title": "What changed since the previous analysis (%s → %s)" % (prev["period"], d["meta"]["period"]),
        "columns": ["Indicator", "%s (%s)" % (prev.get("short_label", "Previous"), prev["period"]), "Current (%s)" % cur, "Trend", "Comment"],
        "widths": [36, 24, 24, 18, 60],
        "rows": [[x["indicator"], x["previous"], x["current"], x["trend"], x.get("comment", "")] for x in prev["rows"]],
        "trend_col": 3,
    }
    ws = detail_sheet(wb, spec)
    r = ws.max_row + 2
    for head, key in (("Closed / no longer reported this cycle", "closed"), ("New this cycle", "new")):
        if prev.get(key):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            ws.cell(r, 1, head).font = font(11, True, NAVY)
            for t in prev[key]:
                r += 1
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
                x = ws.cell(r, 1, "• " + t); x.font = font(10); x.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = row_height([t], [160], 15)
            r += 2


def build(findings_path, out_path):
    d = load_findings(findings_path)
    wb = Workbook()
    exec_summary(wb, d)
    action_items(wb, d)
    for spec in d.get("detail_sheets", []):
        detail_sheet(wb, spec)
    if d.get("previous"):
        change_sheet(wb, d)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    try:
        print("wrote", build(sys.argv[1], sys.argv[2]))
    except FindingsError as e:
        sys.exit(str(e))
